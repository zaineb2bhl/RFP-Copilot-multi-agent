"""Tests for deterministic routing, individual nodes, and the full graph with interrupt."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document
from langgraph.checkpoint.memory import InMemorySaver
from langgraph.types import Command

from rfp_copilot.config import Settings
from rfp_copilot.graph.graph import build_graph
from rfp_copilot.graph.nodes import (
    PipelineDeps,
    apply_deterministic_checks,
    groups_needing_revision,
    make_review_node,
    make_write_node,
    route_after_human,
    route_after_review,
    unsupported_figures,
    write_section,
)
from rfp_copilot.llm import MockLLM
from rfp_copilot.models import Requirement, ReviewResult
from rfp_copilot.retrieval.retriever import HybridRetriever

REQS = [
    Requirement(id="REQ-001", text="Must hold ISO 27001.", type="mandatory", source_page=1),
    Requirement(id="REQ-002", text="Should offer training.", type="desirable", source_page=2),
]


def _review(score: float) -> dict:
    return ReviewResult(coverage=score, groundedness=score).model_dump()


# -- routing ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("scores", "iteration", "expected"),
    [
        ({"G01": 0.9, "G02": 0.8}, 1, "human_review"),  # all pass
        ({"G01": 0.9, "G02": 0.5}, 1, "write"),  # one fails, passes left
        ({"G01": 0.9, "G02": 0.5}, 2, "human_review"),  # one fails, no passes left
        ({}, 1, "human_review"),  # nothing to review
        ({"G01": 0.75}, 1, "human_review"),  # threshold is inclusive
    ],
)
def test_route_after_review(scores: dict, iteration: int, expected: str) -> None:
    state = {
        "reviews": {g: _review(s) for g, s in scores.items()},
        "passages": {g: [{"chunk_id": "kb#1"}] for g in scores},
        "iteration": iteration,
    }
    assert route_after_review(state, threshold=0.75, max_iterations=2) == expected


def test_sections_without_evidence_are_not_rewritten() -> None:
    state = {
        "reviews": {"G01": _review(0.2), "G02": _review(0.2)},
        "passages": {"G01": [{"chunk_id": "kb#1"}], "G02": []},
        "iteration": 1,
    }
    assert groups_needing_revision(state, 0.75) == ["G01"]
    state["passages"]["G01"] = []
    assert route_after_review(state, threshold=0.75, max_iterations=2) == "human_review"


def test_unsupported_figures_are_flagged() -> None:
    passages = {"sla#1": "Critical incidents: response within 1 hour, 24/7. Uptime 99.9%."}
    requirement = "The vendor must respond to critical issues within 2 hours, 24 hours a day."
    good = "- **REQ-001**: We respond to critical incidents within 1 hour, 24/7 [sla#1]."
    assert unsupported_figures(good, passages, requirement) == []
    invented = "- **REQ-001**: We respond within 30 minutes and guarantee 99.99% uptime [sla#1]."
    issues = unsupported_figures(invented, passages, requirement)
    assert len(issues) == 1 and "30" in issues[0] and "99.99%" in issues[0]
    uncited = "- **REQ-002**: We launched 64 city websites."
    assert unsupported_figures(uncited, passages, requirement)
    # The RFP's own figures may be restated without a source.
    assert (
        unsupported_figures("**REQ-001** Within 2 hours: to be confirmed.", {}, requirement) == []
    )
    checked = apply_deterministic_checks(
        ReviewResult(coverage=1, groundedness=1), invented, REQS[:1], passages
    )
    assert checked.groundedness == 0.5 and checked.unsupported_claims


def test_route_after_human() -> None:
    assert route_after_human({"outcome": "approved"}) == "export_docx"
    assert route_after_human({"outcome": "rejected"}) == "__end__"


def test_deterministic_checks_tighten_review() -> None:
    llm_review = ReviewResult(coverage=1.0, groundedness=1.0)
    draft = "**REQ-001** We are certified [company-profile#3] and [made-up#9]."
    checked = apply_deterministic_checks(llm_review, draft, REQS, {"company-profile#3": "x"})
    assert checked.uncovered_requirements == ["REQ-002"]
    assert checked.coverage == 0.5
    assert checked.groundedness == 0.5
    assert any("made-up#9" in c for c in checked.unsupported_claims)


def test_multi_citations_and_heading_cleanup() -> None:
    from rfp_copilot.graph.nodes import citations, clean_draft

    text = "A [x#1, y-2#3]; B [x#1]; C [z#4;w#5] and [not a citation]."
    assert citations(text) == ["x#1", "y-2#3", "z#4", "w#5"]
    draft = "## 5. Security\n\n**REQ-001** ok\n## Sub\ntext"
    assert clean_draft(draft) == "**REQ-001** ok\n## Sub\ntext"


def test_draft_without_citations_is_not_grounded() -> None:
    checked = apply_deterministic_checks(
        ReviewResult(coverage=1, groundedness=1), "REQ-001 REQ-002 all good.", REQS, {}
    )
    assert checked.groundedness == 0.0


def test_administrative_only_sections_need_no_citation() -> None:
    admin = [
        Requirement(id="REQ-009", text="Submit one PDF.", type="administrative", source_page=4)
    ]
    review = ReviewResult(coverage=1, groundedness=0.9)
    checked = apply_deterministic_checks(review, "**REQ-009**: we comply.", admin, {})
    assert checked.groundedness == 0.9


# -- nodes -----------------------------------------------------------------------------


def _state() -> dict:
    return {
        "requirements": [r.model_dump() for r in REQS],
        "groups": [{"id": "G01", "title": "Security", "requirement_ids": ["REQ-001", "REQ-002"]}],
        "passages": {
            "G01": [
                {
                    "chunk_id": "kb#1",
                    "source": "kb.md",
                    "title": "KB",
                    "text": "ISO 27001.",
                    "score": 1.0,
                }
            ]
        },
    }


def _answers(*items: tuple[str, str, str | None]) -> str:
    return json.dumps(
        {
            "answers": [
                {"requirement_id": i, "answer": a, "missing_information": m} for i, a, m in items
            ]
        }
    )


def test_write_node_includes_feedback_on_revision(settings: Settings) -> None:
    first = _answers(("REQ-001", "Certified [kb#1].", None), ("REQ-002", "", "training plan"))
    second = _answers(("REQ-001", "ISO 27001 certified [kb#1].", None), ("REQ-002", "", "plan"))
    writer = MockLLM([first, second])
    deps = PipelineDeps(settings=settings, writer=writer, reviewer=None)
    write = make_write_node(deps)
    state = {**_state(), "iteration": 0}
    update = write(state)
    assert update["drafts"]["G01"] == (
        "- **REQ-001**: Certified [kb#1].\n- **REQ-002**: [TO CONFIRM: training plan]"
    )
    assert update["iteration"] == 1
    state |= update
    state["reviews"] = {"G01": {**_review(0.3), "feedback": "Cite the ISO certificate."}}
    update = write(state)
    assert update["drafts"]["G01"].startswith("- **REQ-001**: ISO 27001 certified [kb#1].")
    revision_prompt = writer.calls[1][-1]["content"]
    assert "Certified [kb#1]." in revision_prompt and "Cite the ISO certificate." in revision_prompt


def test_writer_retries_ranges_and_missing_ids(settings: Settings) -> None:
    grouped = _answers(("REQ-001", "REQ-001 through REQ-002: we comply [kb#1].", None))
    valid = _answers(("REQ-001", "Certified [kb#1].", None), ("REQ-002", "", "training plan"))
    writer = MockLLM([grouped, valid])
    draft = write_section(writer, "Security", REQS, [], issuer="City of Riverton")
    assert draft.count("**REQ-") == 2 and "through" not in draft
    retry_prompt = writer.calls[1][-1]["content"]
    assert "missing answers for REQ-002" in retry_prompt and "range" in retry_prompt
    assert "Buyer: City of Riverton" in writer.calls[0][-1]["content"]


def test_writer_failure_becomes_placeholders_and_admin_skips_llm(settings: Settings) -> None:
    admin = Requirement(id="REQ-009", text="Submit 3 copies.", type="administrative", source_page=4)
    writer = MockLLM(lambda _m: "not json")
    draft = write_section(writer, "Mixed", [REQS[0], admin], [])
    assert "- **REQ-001**: [TO CONFIRM: the answer could not be generated" in draft
    assert "- **REQ-009**: [TO CONFIRM: bid team check: Submit 3 copies]" in draft
    only_admin = MockLLM([])
    assert "bid team check" in write_section(only_admin, "Admin", [admin], [])
    assert only_admin.calls == []  # no LLM call for checklist items


def test_review_node_applies_checks_and_survives_bad_json(settings: Settings) -> None:
    reviewer = MockLLM(lambda _m: "definitely not json")
    deps = PipelineDeps(settings=settings, writer=None, reviewer=reviewer)
    state = {**_state(), "drafts": {"G01": "REQ-001 [kb#1]"}, "pending_groups": ["G01"]}
    update = make_review_node(deps)(state)
    review = ReviewResult.model_validate(update["reviews"]["G01"])
    assert review.score == 0.0
    assert "REQ-002" in review.uncovered_requirements
    assert update["review_history"] == [{"G01": 0.0}]


# -- full graph ------------------------------------------------------------------------


def test_full_graph_loop_interrupt_and_export(
    settings: Settings, sample_pdf: Path, retriever: HybridRetriever, pipeline_llm: MockLLM
) -> None:
    deps = PipelineDeps(
        settings=settings, writer=pipeline_llm, reviewer=pipeline_llm, retriever=retriever
    )
    graph = build_graph(deps, InMemorySaver())
    config = {"configurable": {"thread_id": "t1"}}

    result = graph.invoke({"run_id": "t1", "rfp_path": str(sample_pdf)}, config)
    assert "__interrupt__" in result  # paused for human review
    state = graph.get_state(config)
    assert state.next == ("human_review",)
    values = state.values
    assert values["iteration"] == 2  # reviewer failed every section once -> one revision pass
    assert len(values["review_history"]) == 2
    assert all(score == 0.4 for score in values["review_history"][0].values())
    assert all(score == 0.9 for score in values["review_history"][1].values())
    assert len(values["requirements"]) >= 20

    first_group = values["groups"][0]["id"]
    edited = "Edited by a human: **REQ-001** handled [company-profile#1]."
    result = graph.invoke(
        Command(resume={"action": "approve", "edits": {first_group: edited}}), config
    )
    output = Path(result["output_path"])
    assert output.exists() and result["outcome"] == "approved"
    docx = Document(output)
    text = "\n".join(p.text for p in docx.paragraphs)
    assert "Edited by a human" in text
    assert "Compliance matrix" in text
    table = docx.tables[0]
    assert len(table.rows) == len(values["requirements"]) + 1
    assert table.rows[0].cells[0].text == "Requirement"


def test_full_graph_reject_skips_export(
    settings: Settings, sample_pdf: Path, retriever: HybridRetriever, pipeline_llm: MockLLM
) -> None:
    deps = PipelineDeps(
        settings=settings, writer=pipeline_llm, reviewer=pipeline_llm, retriever=retriever
    )
    graph = build_graph(deps, InMemorySaver())
    config = {"configurable": {"thread_id": "t2"}}
    graph.invoke({"run_id": "t2", "rfp_path": str(sample_pdf)}, config)
    result = graph.invoke(Command(resume={"action": "reject"}), config)
    assert result["outcome"] == "rejected"
    assert "output_path" not in result
    assert not (settings.runs_dir / "outputs").exists()
    json.dumps(result)  # state stays JSON-serialisable for the checkpointer


def test_repeated_requirement_id_is_collapsed() -> None:
    from rfp_copilot.graph.nodes import clean_draft

    assert clean_draft("- **REQ-045**: **REQ-045**: to be confirmed.") == (
        "- **REQ-045**: to be confirmed."
    )


def test_every_requirement_gets_its_own_passages(retriever: HybridRetriever) -> None:
    from rfp_copilot.graph.nodes import section_passages

    queries = {
        "REQ-001": "WCAG 2.1 Level AA accessibility",
        "REQ-002": "vanity URLs and redirect manager",
        "REQ-003": "bids and RFPs module with addenda",
    }
    passages, evidence = section_passages(retriever, queries, budget=6, minimum=2, maximum=10)
    assert len(passages) <= 10
    for req_id, query in queries.items():
        top = retriever.search(query, k=2)[0].chunk_id
        assert top in {p.chunk_id for p in passages}, query
        assert evidence[req_id][0]["chunk_id"] == top
        assert evidence[req_id][0]["similarity"] is not None


def test_figure_check_handles_abbreviations() -> None:
    passages = {"acc#1": "Every website conforms to WCAG 2.1 Level AA under the U.S. DOJ rule."}
    draft = (
        "- **REQ-061**: Every website we build conforms to WCAG 2.1 Level AA, required under the "
        "U.S. Department of Justice rule [acc#1]. We trained 58 editors."
    )
    issues = unsupported_figures(draft, passages, "Demonstrate ADA-compliant web development.")
    assert len(issues) == 1 and "58" in issues[0]  # only the uncited figure, not "2.1"
