"""Tests for DOCX export, the SQLite repository and the FastAPI endpoints."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from rfp_copilot.api.app import create_app
from rfp_copilot.compliance import build_compliance_rows
from rfp_copilot.config import Settings, runs_db_url
from rfp_copilot.db.repository import Database
from rfp_copilot.export.docx_exporter import export_docx
from rfp_copilot.llm import MockLLM
from rfp_copilot.models import (
    NonCompliance,
    Requirement,
    ResponseGroup,
    RetrievedPassage,
    ReviewResult,
)
from rfp_copilot.pipeline import PipelineService
from rfp_copilot.retrieval.retriever import HybridRetriever
from tests.fakes import make_pipeline_llm

REQS = [
    Requirement(id="REQ-001", text="Must hold ISO 27001.", type="mandatory", source_page=1),
    Requirement(id="REQ-002", text="Should offer training.", type="desirable", source_page=2),
    Requirement(id="REQ-003", text="Submit a PDF.", type="administrative", source_page=3),
]
GROUPS = [
    ResponseGroup(id="G01", title="Security", requirement_ids=["REQ-001", "REQ-002", "REQ-003"])
]
DRAFT = "- **REQ-001**: certified [kb#1].\n- **REQ-002**: we train staff [kb#2]."


# -- export ----------------------------------------------------------------------------


RETRIEVED = {"G01": {"kb#1", "kb#2", "kb#3"}}


def test_compliance_rows() -> None:
    draft = (
        "- **REQ-001**: certified [kb#1].\n"
        "- **REQ-002**: we train staff [kb#2]. [TO CONFIRM: number of sessions]\n"
        "- **REQ-003**: [TO CONFIRM: bid team check: Submit a PDF.]"
    )
    rows = {
        r.requirement_id: r
        for r in build_compliance_rows(REQS, GROUPS, {"G01": draft}, {}, RETRIEVED)
    }
    assert (rows["REQ-001"].status, rows["REQ-001"].sources) == ("addressed", ["kb#1"])
    assert rows["REQ-001"].where_addressed == "1.1 Security"
    assert (rows["REQ-002"].status, rows["REQ-002"].note) == ("partial", "number of sessions")
    assert rows["REQ-003"].status == "to_confirm" and rows["REQ-003"].sources == []
    assert rows["REQ-003"].note.startswith("Bid team check: Submit a PDF.")


def test_non_compliance_is_flagged_with_note() -> None:
    review = ReviewResult(
        coverage=1,
        groundedness=1,
        non_compliant=[
            NonCompliance(requirement_id="REQ-001", required="ISO 27001", offered="none")
        ],
    )
    rows = build_compliance_rows(REQS, GROUPS, {"G01": DRAFT}, {"G01": review}, RETRIEVED)
    assert rows[0].status == "non_compliant" and "ISO 27001" in rows[0].note


def test_export_docx_content(tmp_path: Path) -> None:
    review = ReviewResult(
        coverage=1,
        groundedness=1,
        non_compliant=[
            NonCompliance(requirement_id="REQ-002", required="5 days", offered="2 days")
        ],
    )
    rows = build_compliance_rows(REQS, GROUPS, {"G01": DRAFT}, {"G01": review}, RETRIEVED)
    passages = {
        "kb#1": RetrievedPassage(chunk_id="kb#1", source="kb.md", title="KB", text="t", score=1)
    }
    path = export_docx(
        tmp_path / "out" / "r.docx",
        rfp_title="Test RFP",
        groups=GROUPS,
        drafts={"G01": DRAFT},
        compliance=rows,
        passages=passages,
    )
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Technical Response: Test RFP" in text
    assert "Warnings for the bid team" in text and "5 days" in text
    assert "REQ-001: certified [kb#1]." in text  # bold markers rendered, not printed
    assert "[kb#1] kb.md, KB" in text
    assert "3. Bid team checklist" in text and "☐ REQ-003" in text
    bold_runs = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "REQ-001" in bold_runs
    cells = [[c.text for c in row.cells] for row in doc.tables[0].rows]
    assert cells[0] == ["Requirement", "Status", "Where addressed", "Sources", "Notes"]
    assert [c[1] for c in cells[1:]] == ["Addressed", "NON-COMPLIANT", "To confirm"]


# -- database --------------------------------------------------------------------------


def test_database_roundtrip(tmp_path: Path) -> None:
    db = Database(f"sqlite:///{(tmp_path / 'db' / 'app.sqlite').as_posix()}")
    db.create_run("r1", "rfp.pdf")
    db.create_run("r2", "other.pdf")
    run = db.update_run("r1", status="running", stage="ingest", usage={"calls": 3})
    assert (run.status, run.stage, run.usage) == ("running", "ingest", {"calls": 3})
    assert [r.id for r in db.list_runs()] == ["r2", "r1"]
    rows = build_compliance_rows(REQS, GROUPS, {"G01": DRAFT}, {}, RETRIEVED)
    db.replace_requirements("r1", REQS, rows)
    stored = db.get_requirements("r1")
    assert [r.status for r in stored] == ["addressed", "addressed", "to_confirm"]
    db.replace_requirements("r1", REQS[:1])
    assert [r.status for r in db.get_requirements("r1")] == ["pending"]
    assert db.get_run("missing") is None
    with pytest.raises(KeyError):
        db.update_run("missing", status="failed")


# -- API -------------------------------------------------------------------------------


@pytest.fixture
def client(settings: Settings, retriever: HybridRetriever) -> TestClient:
    handler = make_pipeline_llm()
    service = PipelineService(
        settings,
        llm_factory=lambda role, usage: MockLLM(handler, usage=usage),
        retriever=retriever,
    )
    return TestClient(create_app(service))


def test_api_full_flow(client: TestClient, sample_pdf: Path) -> None:
    assert client.get("/health").json()["status"] == "ok"
    # TestClient runs background tasks before returning, so the run is already paused.
    response = client.post(
        "/runs", files={"file": ("rfp.pdf", sample_pdf.read_bytes(), "application/pdf")}
    )
    assert response.status_code == 202
    run_id = response.json()["id"]
    run = client.get(f"/runs/{run_id}").json()
    assert run["status"] == "awaiting_review", run["error"]
    assert run["usage"]["calls"] > 0

    requirements = client.get(f"/runs/{run_id}/requirements").json()
    assert requirements and all(r["status"] == "pending" for r in requirements)
    review = client.get(f"/runs/{run_id}/review").json()
    assert review["iteration"] == 2 and review["groups"]
    assert client.get(f"/runs/{run_id}/download").status_code == 404

    response = client.post(f"/runs/{run_id}/review", json={"action": "approve", "edits": {}})
    assert response.status_code == 202
    run = client.get(f"/runs/{run_id}").json()
    assert run["status"] == "completed", run["error"]
    statuses = {r["status"] for r in client.get(f"/runs/{run_id}/requirements").json()}
    assert statuses <= {"addressed", "partial", "not_covered", "to_confirm", "non_compliant"}
    download = client.get(f"/runs/{run_id}/download")
    assert download.status_code == 200 and download.content[:2] == b"PK"
    assert [r["id"] for r in client.get("/runs").json()] == [run_id]
    assert client.post(f"/runs/{run_id}/review", json={"action": "approve"}).status_code == 409


def test_api_rejects_non_pdf_and_unknown_run(client: TestClient) -> None:
    response = client.post("/runs", files={"file": ("notes.txt", b"hello", "text/plain")})
    assert response.status_code == 415
    response = client.post("/runs", files={"file": ("fake.pdf", b"hello", "application/pdf")})
    assert response.status_code == 415
    assert client.get("/runs/nope").status_code == 404


def test_failed_run_is_reported(
    settings: Settings, retriever: HybridRetriever, tmp_path: Path
) -> None:
    service = PipelineService(
        settings,
        llm_factory=lambda role, usage: MockLLM(make_pipeline_llm(), usage=usage),
        retriever=retriever,
    )
    run = service.create_run("broken.pdf", b"%PDF-1.4 not really a pdf")
    run = service.execute(run.id)
    assert run.status == "failed" and run.error
    assert Database(runs_db_url(settings)).get_run(run.id).status == "failed"
